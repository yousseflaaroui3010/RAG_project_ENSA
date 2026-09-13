"""ST-13 exit gate: the fixture corpus converts, a corrupted file reports
Failed, and a scanned PDF reports Skipped with a reason.

Reference: docs/phase2/Sanad_Architecture_v1.0.md ADR-07 (the ladder:
pymupdf4llm for PDF, markitdown for DOCX, passthrough for TXT and MD),
docs/phase2/Sanad_PRD_v1.0.md F-02 (criterion 3: one failing file never
blocks the rest), F-16 (binding V1: a scanned PDF is Skipped with the
reason stated) and section 11's failure table.

Every fixture is a real file with real bytes, built here rather than
committed: a genuine multi-page PDF from pymupdf, a genuine OOXML package
from `zipfile`, a genuinely encrypted PDF, genuinely corrupted bytes. The
converters are never mocked. That is the point -- ST-13's whole risk lives
in what two third-party parsers do with hostile input, and a mocked parser
would answer that question by assuming it. The ST-11 review found a test
that passed while proving nothing; a fake PDF would be that test again.

Assertions come in pairs wherever a mutation could restore a DEFAULT
rather than remove behaviour (the ST-12 lesson recorded in BUILD-STATE):
a one-sided bound cannot see a value that falls back to something
plausible.
"""

from __future__ import annotations

import logging
import re
import zipfile
from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation

import chunking
import conversion
from config import Settings, get_settings
from conversion import ConversionOutcome

# --- fixture builders -------------------------------------------------


def _pdf_with_headings(path, pages=2):
    """A real PDF whose big text becomes markdown headings. Two pages by
    default so a page count of 1 (a plausible hardcoded default) is
    distinguishable from a page count that was actually read."""
    doc = pymupdf.open()
    for number in range(1, pages + 1):
        page = doc.new_page()
        page.insert_text((72, 100), f"Chapitre {number}", fontsize=24)
        page.insert_text((72, 140), f"Corps du texte de la page {number}.", fontsize=11)
    doc.save(path)
    doc.close()
    return path


def _scanned_pdf(path):
    """A PDF with an image and no text layer at all -- what a flatbed
    scanner produces, and the F-16 binding case."""
    doc = pymupdf.open()
    page = doc.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200))
    pixmap.set_rect(pixmap.irect, (255, 240, 200))
    page.insert_image(pymupdf.Rect(50, 50, 250, 250), pixmap=pixmap)
    doc.save(path)
    doc.close()
    return path


def _locked_pdf(path):
    """A real AES-256 encrypted PDF. Note it OPENS without error; only
    `needs_pass` reveals it, which is exactly the trap `_read_pdf` guards."""
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Confidentiel", fontsize=20)
    doc.save(
        path,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="user",
    )
    doc.close()
    return path


_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml"
 ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1"
 Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
 Target="word/document.xml"/>
</Relationships>"""

_DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Titre principal</w:t></w:r></w:p>
<w:p><w:r><w:t>Un paragraphe de contenu.</w:t></w:r></w:p>
<w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>Sous titre</w:t></w:r></w:p>
</w:body>
</w:document>"""


def _docx(path, document_xml=_DOCUMENT_XML):
    """A genuine minimal OOXML package, assembled with the standard
    library. Built by hand on purpose: python-docx is not a dependency of
    this project, and adding one to write test fixtures would be a real
    dependency bought with a test's convenience."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", _CONTENT_TYPES)
        package.writestr("_rels/.rels", _RELS)
        package.writestr("word/document.xml", document_xml)
    return path


def _pptx(path, slides, *, titles=None):
    """A genuine PPTX package, built with python-pptx -- a transitive
    dependency of `markitdown[all]` (already a project dependency, ADR-07),
    used directly here the same way `pymupdf` is used above to build real
    PDF fixtures: the risk this rung carries lives in what a real parser
    does with a real file, and a mocked one would answer that by assuming
    it. Not a project dependency in its own right (see DECISIONS.md).

    `slides` is one entry per slide: a body string for a slide with text,
    or None for a genuinely blank slide -- F-11's binding empty-slide case,
    and the one a naive "count slides seen so far" numbering gets wrong.
    `titles`, if given, is the same length and sets a real title
    placeholder on the matching slide (None entries skip it)."""
    prs = Presentation()
    layout_with_body = prs.slide_layouts[1]
    layout_blank = prs.slide_layouts[6]
    titles = titles or [None] * len(slides)
    for text, title in zip(slides, titles, strict=True):
        if text is None:
            prs.slides.add_slide(layout_blank)
            continue
        slide = prs.slides.add_slide(layout_with_body)
        if title is not None:
            slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = text
    prs.save(str(path))
    return path


# --- the happy rungs --------------------------------------------------


def test_pdf_converts_and_preserves_headings(tmp_path):
    """ADR-07's entire reason for keeping PDFs off markitdown: the parent
    splitter cuts on markdown headings (architecture 7.5), so a converter
    that flattens them destroys the structure retrieval depends on.

    Both halves are asserted. Heading markup alone would still pass if the
    converter dropped the body; body text alone would still pass if it
    dropped every heading -- which is precisely the regression ADR-07 is
    written to prevent."""
    result = conversion.convert_file(_pdf_with_headings(tmp_path / "code.pdf"))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.reason is None
    assert "# Chapitre 1" in result.markdown
    assert "# Chapitre 2" in result.markdown
    assert "Corps du texte de la page 1." in result.markdown
    assert "Corps du texte de la page 2." in result.markdown


def test_pdf_reports_its_real_page_count(tmp_path):
    """PRD F-02 criterion 1 promises each added file is reported "with its
    page count", and db/schema.sql has the column for it.

    Three page counts are checked, not one. A single assertion against a
    2-page file cannot tell a real read from a hardcoded 2, and `None` or
    `1` are both plausible defaults a mutation could restore."""
    counts = {}
    for pages in (1, 2, 5):
        result = conversion.convert_file(
            _pdf_with_headings(tmp_path / f"doc{pages}.pdf", pages=pages)
        )
        assert result.outcome is ConversionOutcome.CONVERTED
        counts[pages] = result.page_count

    assert counts == {1: 1, 2: 2, 5: 5}


def test_docx_converts_and_preserves_headings(tmp_path):
    result = conversion.convert_file(_docx(tmp_path / "note.docx"))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert "# Titre principal" in result.markdown
    assert "## Sous titre" in result.markdown
    assert "Un paragraphe de contenu." in result.markdown


def test_docx_has_no_page_count(tmp_path):
    """None means "this format has no pages", not "we did not look". A
    DOCX has no fixed pagination until something renders it, so any number
    here would be invented."""
    result = conversion.convert_file(_docx(tmp_path / "note.docx"))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.page_count is None


def test_pptx_converts_and_labels_each_slide(tmp_path):
    """F-11's core claim, at the conversion layer: markitdown's own
    `<!-- Slide number: N -->` marker (verified against a real deck)
    becomes a real "Slide N" heading, one per slide with text.

    Slide 2 is genuinely blank -- the middle slide, not the last, so a
    numbering scheme that merely counts headings seen so far would land on
    2 for slide 3's text instead of 3. Asserted by absence too: "Slide 2"
    must not appear anywhere, or a blank slide would be inventing a
    citable but empty section."""
    path = _pptx(
        tmp_path / "deck.pptx",
        ["Alpha content on slide one.", None, "Gamma content on slide three."],
    )

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert "# Slide 1" in result.markdown
    assert "# Slide 3" in result.markdown
    assert "Slide 2" not in result.markdown
    assert "Alpha content on slide one." in result.markdown
    assert "Gamma content on slide three." in result.markdown


def test_pptx_slide_title_is_kept_as_text_not_a_second_heading(tmp_path):
    """markitdown emits a slide's own title as its own "# Title" heading.
    Left alone, chunking would split that into a second section with no
    slide number at all -- so `_read_pptx` demotes it to plain text, and
    this is the one test that would catch that demotion being dropped:
    exactly one H1 for the slide, and the title text still present, just
    not as a heading."""
    path = _pptx(tmp_path / "deck.pptx", ["Alpha content."], titles=["Intro"])

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    heading_lines = [
        line for line in result.markdown.splitlines() if line.startswith("#")
    ]
    assert heading_lines == ["# Slide 1"]
    assert "Intro" in result.markdown


def test_pptx_has_no_page_count(tmp_path):
    """None means "this format has no pages", matching DOCX: a slide deck
    is cited by slide number (PRD F-11), never a page."""
    result = conversion.convert_file(_pptx(tmp_path / "deck.pptx", ["Some text."]))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.page_count is None


def test_pptx_with_no_text_reports_skipped(tmp_path):
    """A deck of genuinely blank slides reaches the same shared emptiness
    gate as a scanned PDF or an empty DOCX -- nothing to index, and not the
    user's fault."""
    path = _pptx(tmp_path / "deck.pptx", [None, None])

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.SKIPPED
    assert result.markdown is None


def test_corrupted_pptx_reports_failed_and_never_returns_its_own_bytes(tmp_path):
    """The PPTX twin of `_read_docx`'s most dangerous behaviour, verified
    against markitdown 0.1.5 rather than assumed: a `.pptx` that is not
    even a ZIP comes back as its own raw bytes, reported a SUCCESS, with no
    exception at all. Reported Added, those bytes would be chunked,
    embedded, and one day cited to a user as a source.

    The second assertion is the load-bearing one, for the same reason it
    is in the DOCX test: FAILED alone would still pass if the guard were
    removed and something else happened to reject the file; garbage can
    never be in `markdown` while the guard holds."""
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"PK\x03\x04garbage-not-a-zip")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None
    assert "PowerPoint" in result.reason


def test_valid_zip_that_is_not_a_pptx_file_reports_failed(tmp_path):
    """The half of the fallback a bare `zipfile.is_zipfile` check would
    miss: a real ZIP renamed `.pptx` passes the zip check, and markitdown
    returns a directory listing of its entries as the document's text.
    Only the `ppt/presentation.xml` check catches it."""
    path = tmp_path / "archive.pptx"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("hello.txt", "hi")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None


def test_pptx_with_a_corrupt_body_reports_failed(tmp_path):
    """The rung's last guard: a package that has `ppt/presentation.xml` by
    name but whose content is not something python-pptx can open (missing
    the rest of the OOXML package here, invalid XML in the DOCX sibling
    test). Passes both the zip and the main-part checks, so only
    markitdown's own `MarkItDownException` catches it -- verified
    reachable against a real minimal fixture, not assumed."""
    path = tmp_path / "mangled.pptx"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("ppt/presentation.xml", "not xml at all <<<>>>")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None
    assert "PowerPoint" in result.reason


def test_pptx_chunking_labels_the_right_slide_not_the_count_so_far(tmp_path, monkeypatch):
    """The end-to-end F-11 claim: real conversion feeds real chunking with
    no code change to chunking.py at all, and the citation a user would see
    names the ACTUAL slide, not the position among slides that had text.

    `parent_merge_below_chars` is forced to 1 (each section already clears
    it) so slides 1 and 3 land in separate parents instead of merging --
    the merged case, and its range label, is covered separately below."""
    path = _pptx(
        tmp_path / "deck.pptx",
        ["Alpha content on slide one.", None, "Gamma content on slide three."],
    )
    settings = get_settings().model_copy(update={"parent_merge_below_chars": 1})
    monkeypatch.setattr(chunking, "get_settings", lambda: settings)

    result = conversion.convert_file(path)
    chunked = chunking.chunk_document(result.markdown, source_file="deck.pptx")

    assert [parent.section_label for parent in chunked.parents] == ["Slide 1", "Slide 3"]
    slide_three_child = next(c for c in chunked.children if "Gamma" in c.text)
    assert slide_three_child.section_label == "Slide 3"


def test_pptx_chunk_spanning_slides_gets_a_range_label(tmp_path):
    """Consistent with the existing merged-range convention chunking.py
    already uses for headings and Article markers (`_merged_label` /
    `_citation_label`'s "Article 1 ... Article 3"): a parent built from
    several short slides is cited by the range it spans, never by its
    first slide alone -- citing a sentence from slide 3 as "Slide 1" would
    be as wrong a citation as citing Article 3 as Article 1.

    Default settings on purpose (no monkeypatch): short real slide text
    merges under the real `parent_merge_below_chars` (2,000), which is
    exactly the input this behaviour is for."""
    path = _pptx(tmp_path / "deck.pptx", ["Alpha.", None, "Gamma, distinct."])

    result = conversion.convert_file(path)
    chunked = chunking.chunk_document(result.markdown, source_file="deck.pptx")

    assert len(chunked.parents) == 1
    assert chunked.parents[0].section_label == "Slide 1 ... Slide 3"


def test_each_pptx_child_is_cited_by_exactly_the_slides_its_text_comes_from(tmp_path):
    """The citation a user reads comes from the CHILD. With default settings
    a deck of short slides merges into one parent ("Slide 1 ... Slide 5"),
    and before the in-text slide lines every child inherited that whole
    range. Each slide's words carry its own prefix here, so the slides a
    child really holds can be read back from its text and compared with its
    label: no child may be cited by a slide it holds nothing from, and none
    may omit one it does. Slide 2 is blank, so the numbers are not 1..4.

    Kills: dropping "Slide" from the citation marker pattern (every child
    gets the parent range); dropping the closing "(end of Slide N)" line
    (the child that runs from slide 3 into slide 4 is cited "Slide 4")."""
    words = {"alpha": 1, "gamma": 3, "delta": 4, "omega": 5}

    def slide_text(word: str) -> str:
        return " ".join(f"{word}{i}" for i in range(40)) + "."

    path = _pptx(
        tmp_path / "deck.pptx",
        [slide_text("alpha"), None, slide_text("gamma"), slide_text("delta"), slide_text("omega")],
    )

    chunked = chunking.chunk_document(
        conversion.convert_file(path).markdown, source_file="deck.pptx"
    )

    assert [p.section_label for p in chunked.parents] == ["Slide 1 ... Slide 5"]
    assert len(chunked.children) >= 3, "fixture too small: children must span slides"
    for child in chunked.children:
        held = sorted({words[w] for w in re.findall(r"(alpha|gamma|delta|omega)\d", child.text)})
        expected = (
            f"Slide {held[0]}"
            if held[0] == held[-1]
            else f"Slide {held[0]} ... Slide {held[-1]}"
        )
        assert child.section_label == expected, child.text[:60]


def test_an_empty_slide_title_leaves_no_stray_hash_in_the_text(tmp_path):
    """markitdown writes an empty title placeholder as a bare "#" line
    (verified on a real deck). It must be demoted like any other in-slide
    heading, not stored and shown as a lone "#"."""
    path = _pptx(tmp_path / "deck.pptx", ["Body text only, no title."])

    markdown = conversion.convert_file(path).markdown

    assert "Body text only, no title." in markdown
    assert [line for line in markdown.splitlines() if line.strip() == "#"] == []


def test_markdown_file_passes_through_unchanged(tmp_path):
    source = "# Titre\n\nUn paragraphe avec des accents: congés payés.\n"
    path = tmp_path / "readme.md"
    path.write_text(source, encoding="utf-8")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.markdown == source
    assert result.page_count is None


def test_txt_passes_through_unchanged(tmp_path):
    source = "Ceci est un document texte simple.\nDeuxieme ligne.\n"
    path = tmp_path / "notes.txt"
    path.write_text(source, encoding="utf-8")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.markdown == source


def test_utf8_bom_is_stripped_from_the_first_heading(tmp_path):
    """Windows editors prepend a byte-order mark. Decoded as plain UTF-8 it
    survives as an invisible first character, so the file starts with
    '\\ufeff# Titre' and the chunker's heading match fails on a document
    that looks perfect in every editor. `text_file_encoding` is
    "utf-8-sig" for this reason.

    Asserted from both ends: the BOM is gone AND the heading is intact at
    position zero. Checking only that the BOM is absent would also pass if
    the whole first line had been eaten."""
    path = tmp_path / "bom.md"
    path.write_bytes("# Titre\n\nCorps.\n".encode("utf-8-sig"))

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert "﻿" not in result.markdown
    assert result.markdown.startswith("# Titre")


# --- the Failed rung (PRD F-02 criterion 3, section 11 failure table) --


def test_corrupted_pdf_reports_failed_with_a_reason(tmp_path):
    """Exit gate. The reason is user-facing (it lands in sync_item.reason
    and renders in the S2 file table), so it is asserted to be real prose,
    not merely non-None."""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.7\nthis is not a pdf at all\n" + b"\x00\xff" * 50)

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None
    assert "damaged" in result.reason
    assert len(result.reason.split()) > 3


def test_empty_pdf_reports_failed(tmp_path):
    """A zero-byte .pdf raises EmptyFileError, a SUBCLASS of FileDataError.
    Held here so a future narrowing of that except clause to the subclass
    list cannot silently drop the truncated-download case."""
    path = tmp_path / "empty.pdf"
    path.write_bytes(b"")

    assert conversion.convert_file(path).outcome is ConversionOutcome.FAILED


def test_password_protected_pdf_reports_failed_with_a_reason(tmp_path):
    """PRD section 11: "Corrupted or password-protected file -> Failed with
    reason". The trap this covers: pymupdf OPENS an encrypted PDF without
    complaint and even answers `page_count`; only `needs_pass` tells the
    truth. Without that check pymupdf4llm dies on an undeclared TypeError
    mid-batch.

    Asserted as FAILED *and* on the password wording, because a generic
    failure would technically satisfy the first half while telling the
    user nothing they can act on."""
    result = conversion.convert_file(_locked_pdf(tmp_path / "locked.pdf"))

    assert result.outcome is ConversionOutcome.FAILED
    assert "password" in result.reason


def test_corrupted_docx_reports_failed_and_never_returns_its_own_bytes(tmp_path):
    """The most dangerous behaviour in this module, and the reason
    `_read_docx` validates the package itself.

    markitdown 0.1.5 does not raise on a `.docx` it cannot parse as Word:
    it falls through to another converter and returns the raw bytes AS A
    SUCCESS. Verified -- this exact fixture came back as the string
    'PK\\x03\\x04garbage-not-a-zip' with no exception. Reported Added,
    those bytes would be chunked, embedded, and one day cited to a user as
    a source.

    So the second assertion is the load-bearing one: FAILED alone would
    still pass if the guard were removed and something else happened to
    reject it, but garbage can never be in `markdown` while this holds."""
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04garbage-not-a-zip")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None
    assert "Word" in result.reason


def test_valid_zip_that_is_not_a_word_file_reports_failed(tmp_path):
    """The half of the fallback that `zipfile.is_zipfile` alone misses: a
    real ZIP renamed .docx passes the zip check, and markitdown returns a
    directory listing of its entries as the document's text. Only the
    `word/document.xml` check catches it."""
    path = tmp_path / "archive.docx"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("hello.txt", "hi")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None


def test_docx_with_a_corrupt_body_reports_failed(tmp_path):
    """The rung's last guard: a package that IS a Word file by structure
    but whose `word/document.xml` is not valid XML. It passes both the zip
    and the main-part checks, so only markitdown's own
    `MarkItDownException` catches it -- verified reachable, not assumed."""
    path = _docx(tmp_path / "mangled.docx", document_xml="not xml at all <<<>>>")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert result.markdown is None
    assert "Word" in result.reason


def test_pdf_renamed_to_docx_reports_failed(tmp_path):
    """A PDF with the wrong extension must not quietly take the DOCX rung:
    markitdown converts it with its own PDF path, which ADR-07 rejects
    because it strips heading structure. The user gets one honest Failed
    row instead of a silently degraded document."""
    source = _pdf_with_headings(tmp_path / "real.pdf")
    path = tmp_path / "mislabelled.docx"
    path.write_bytes(source.read_bytes())

    assert conversion.convert_file(path).outcome is ConversionOutcome.FAILED


def test_undecodable_text_file_reports_failed(tmp_path):
    """Decoding is strict on purpose. A latin-1 fallback never raises, so
    a mis-encoded French file would be indexed as 'CongÃ©s payÃ©s' and only
    surface when an answer quoted the mojibake back."""
    path = tmp_path / "legacy.txt"
    path.write_bytes(b"Cong\xe9s pay\xe9s\xff\xfe invalid \x80\x81")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert "UTF-8" in result.reason


@pytest.mark.parametrize("name", ["gone.pdf", "gone.docx", "gone.pptx", "gone.txt"])
def test_an_ordinary_missing_file_is_not_logged_as_an_unexpected_failure(
    tmp_path, caplog, name
):
    """.claude/rules/backend.md: "Keep logs free of noise so real bugs
    stand out like a spraying leak, never a slow rot."

    A file deleted mid-sync is routine. If it reached the last-resort
    catch it would be logged ERROR with a full traceback on every sync,
    and the one log line that means "a parser crashed in a way we have
    never seen" would be buried in false alarms. Each rung therefore
    names this case explicitly, and this test is what stops a future
    simplification from quietly merging them back.

    Asserted against the paired test below, which proves a genuinely
    undeclared crash IS logged -- silence alone would also be produced by
    a logger that never fires at all."""
    with caplog.at_level(logging.ERROR, logger="conversion"):
        result = conversion.convert_file(tmp_path / name)

    assert result.outcome is ConversionOutcome.FAILED
    assert caplog.records == []


def test_a_converter_crashing_in_an_undeclared_way_is_logged_with_a_traceback(
    monkeypatch, tmp_path, caplog
):
    """The other bound. Containing a surprise must never mean hiding it
    (.claude/rules/backend.md: "Never catch an error without logging the
    full traceback")."""

    def explode(path):
        raise RuntimeError("undeclared parser bug")

    monkeypatch.setitem(conversion._CONVERTERS, "txt", explode)
    path = tmp_path / "landmine.txt"
    path.write_text("content", encoding="utf-8")

    with caplog.at_level(logging.ERROR, logger="conversion"):
        conversion.convert_file(path)

    assert len(caplog.records) == 1
    assert caplog.records[0].exc_info is not None


def test_a_converter_crashing_in_an_undeclared_way_costs_one_file(monkeypatch, tmp_path):
    """The last-resort catch. Both parsers have already been caught
    throwing something undeclared on hostile input, and F-02 criterion 3
    says one surprise costs one file, never the batch.

    The failure is injected rather than found, because a bug we can
    reproduce on demand is one we have already fixed."""

    def explode(path):
        raise RuntimeError("undeclared parser bug")

    monkeypatch.setitem(conversion._CONVERTERS, "txt", explode)
    path = tmp_path / "landmine.txt"
    path.write_text("content", encoding="utf-8")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.FAILED
    assert "undeclared parser bug" in result.reason


@pytest.mark.parametrize(
    "name", ["gone.txt", "gone.md", "gone.pdf", "gone.docx", "gone.pptx"]
)
def test_missing_file_reports_failed_on_every_rung(tmp_path, name):
    """A file deleted between the scan and the conversion is an ordinary
    race in a sync, and every rung has to survive it identically.

    Parametrised across all five because each rung reaches it by a
    different exception: OSError for text, `pymupdf.FileNotFoundError`
    (which shadows the builtin and subclasses RuntimeError, so `except
    OSError` does NOT catch it) for PDF, and the builtin OSError for
    DOCX and PPTX (both raised by `zipfile.ZipFile` on a missing path).
    The reason is asserted too -- a missing file must not be described as
    damaged, or its owner goes hunting for corruption in a file that
    simply is not there."""
    result = conversion.convert_file(tmp_path / name)

    assert result.outcome is ConversionOutcome.FAILED
    assert "could not be read" in result.reason
    assert "damaged" not in result.reason


# --- the Skipped rung (F-16 binding V1 behaviour) ---------------------


def test_scanned_pdf_reports_skipped_with_a_reason(tmp_path):
    """Exit gate, and the binding V1 behaviour under PRD F-16: a scanned
    PDF is Skipped WITH THE REASON STATED, not Failed -- nothing is
    broken and there is nothing for the user to fix.

    The wording is asserted because "with the reason stated" is the
    requirement; a bare Skipped with a null reason satisfies the enum and
    fails the promise."""
    result = conversion.convert_file(_scanned_pdf(tmp_path / "scan.pdf"))

    assert result.outcome is ConversionOutcome.SKIPPED
    assert result.markdown is None
    assert "no text layer" in result.reason
    assert "scan" in result.reason


def test_scanned_pdf_is_not_reported_as_failed(tmp_path):
    """Guards the distinction itself. Skipped and Failed are different
    promises to the user -- "nothing to index here" versus "this file is
    broken, go fix it" -- and collapsing them would send someone hunting
    for corruption in a perfectly good scan."""
    assert conversion.convert_file(_scanned_pdf(tmp_path / "scan.pdf")).outcome is not (
        ConversionOutcome.FAILED
    )


def test_empty_text_file_reports_skipped_without_the_scan_wording(tmp_path):
    """An empty .txt has nothing to index either, but telling its owner
    the file "is a scan or images only" would be nonsense. Same outcome,
    honest reason."""
    path = tmp_path / "blank.txt"
    path.write_text("   \n\n  \t\n", encoding="utf-8")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.SKIPPED
    assert "no text" in result.reason
    assert "scan" not in result.reason


def test_docx_with_no_text_reports_skipped(tmp_path):
    """A valid Word file of empty paragraphs reaches the same shared
    emptiness gate as a scan."""
    empty_body = _DOCUMENT_XML.replace("Titre principal", "").replace(
        "Un paragraphe de contenu.", ""
    ).replace("Sous titre", "")
    path = _docx(tmp_path / "blank.docx", document_xml=empty_body)

    assert conversion.convert_file(path).outcome is ConversionOutcome.SKIPPED


def test_unsupported_extension_reports_skipped_with_the_shared_reason(tmp_path):
    """A file type outside the V1.1 set (PPTX joined it for F-11 / ST-48;
    RTF has no story of its own). The reason is the SAME constant the
    folder scan uses -- asserted by identity, not by repeating the string
    here, so one reworded copy cannot drift away from the other and show
    the user two explanations for one situation."""
    path = tmp_path / "memo.rtf"
    path.write_bytes(b"anything")

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.SKIPPED
    assert result.reason == conversion.UNSUPPORTED_TYPE_REASON


def test_min_text_chars_is_read_from_config_at_both_bounds(tmp_path, monkeypatch):
    """The ST-12 lesson, applied: a mutation that ignores this setting and
    hardcodes 1 restores a DEFAULT rather than removing behaviour, so a
    test run at the default value cannot see it.

    Run at a NON-default threshold, and asserted on both sides of it:
    just below must Skip, just above must Convert. Either assertion alone
    is satisfiable by a converter that always answers the same way."""
    settings = get_settings().model_copy(update={"conversion_min_text_chars": 10})
    monkeypatch.setattr(conversion, "get_settings", lambda: settings)

    below = tmp_path / "short.txt"
    below.write_text("123456789", encoding="utf-8")  # 9 chars
    above = tmp_path / "long.txt"
    above.write_text("12345678901", encoding="utf-8")  # 11 chars

    assert conversion.convert_file(below).outcome is ConversionOutcome.SKIPPED
    assert conversion.convert_file(above).outcome is ConversionOutcome.CONVERTED


# --- OCR (F-16): the scanned-PDF rung, off by default -----------------
#
# `_ocr_page` is the one small function `_read_pdf` calls to OCR a page
# (see conversion.py). Every test below except the "real" ones fakes it
# with `monkeypatch.setattr(conversion, "_ocr_page", ...)`, so the OCR
# PLUMBING (page mapping, the page cap, error handling, the reason wording)
# is proven without needing the real tessdata files this machine happens
# to have. The "real" tests at the end call the true PyMuPDF/Tesseract
# path and are skipped where the tessdata folder is absent (CI).


def _ocr_settings(**overrides):
    """A Settings copy with OCR turned on. `model_copy` does not
    re-validate (same trick `test_min_text_chars_is_read_from_config_at_
    both_bounds` already relies on above), so the fake tessdata path below
    never has to exist on disk for the plumbing tests."""
    return get_settings().model_copy(
        update={"ocr_tessdata_dir": "Z:/fake-tessdata-not-on-disk", **overrides}
    )


def _scanned_pdf_pages(path, count):
    """A PDF with `count` pages, each an image and no text layer at all --
    the multi-page twin of `_scanned_pdf` above, used where a single page
    cannot distinguish a correct OCR page-mapping from a wrong one."""
    doc = pymupdf.open()
    for _ in range(count):
        page = doc.new_page()
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50))
        pixmap.set_rect(pixmap.irect, (255, 240, 200))
        page.insert_image(pymupdf.Rect(10, 10, 60, 60), pixmap=pixmap)
    doc.save(path)
    doc.close()
    return path


def test_ocr_off_by_default_scanned_pdf_is_skipped_exactly_as_before(
    tmp_path, monkeypatch
):
    """The default (`ocr_tessdata_dir=""`) must leave V1's binding
    behaviour exactly alone. Proven two ways at once: the outcome AND
    wording are unchanged, and `_ocr_page` is never even called -- so a
    future bug that starts OCR'ing on an unconfigured machine fails this
    test even if it happened to still land on SKIPPED."""

    def must_not_run(*args, **kwargs):
        raise AssertionError("_ocr_page must not run when OCR is off")

    monkeypatch.setattr(conversion, "_ocr_page", must_not_run)

    result = conversion.convert_file(_scanned_pdf(tmp_path / "scan.pdf"))

    assert result.outcome is ConversionOutcome.SKIPPED
    assert "no text layer" in result.reason


def test_ocr_on_does_not_touch_a_pdf_that_already_has_a_text_layer(
    tmp_path, monkeypatch
):
    """OCR being configured must never touch a PDF that already converted
    fine: the trigger is "produced no text at all", not "OCR is on", so an
    ordinary typed PDF keeps pymupdf4llm's heading-preserving markdown
    untouched rather than being overwritten by an OCR pass over its
    rendered pixels. This is also the documented mixed-PDF limitation:
    a document with SOME real text stays on this path entirely."""
    monkeypatch.setattr(conversion, "get_settings", lambda: _ocr_settings())

    def must_not_run(*args, **kwargs):
        raise AssertionError("_ocr_page must not run on a PDF with a text layer")

    monkeypatch.setattr(conversion, "_ocr_page", must_not_run)

    result = conversion.convert_file(_pdf_with_headings(tmp_path / "typed.pdf"))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert "Chapitre 1" in result.markdown


def test_ocr_on_converts_a_scanned_pdf_and_the_text_becomes_citable(
    tmp_path, monkeypatch
):
    """The acceptance criterion end to end: OCR text flows through the
    SAME citation-marker regex every other PDF uses (config's
    `parent_citation_marker_pattern`), with no PDF-specific change to
    chunking.py at all."""
    monkeypatch.setattr(conversion, "get_settings", lambda: _ocr_settings())
    monkeypatch.setattr(
        conversion,
        "_ocr_page",
        lambda page, **kwargs: (
            "Article 12: Le salarie a droit a un conge annuel paye."
        ),
    )

    result = conversion.convert_file(_scanned_pdf(tmp_path / "scan.pdf"))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.reason is None
    assert "Article 12" in result.markdown
    assert result.page_count == 1

    chunked = chunking.chunk_document(result.markdown, source_file="scan.pdf")
    assert any(parent.section_label == "Article 12" for parent in chunked.parents)


def test_ocr_page_mapping_is_correct_not_reversed_or_off_by_one(
    tmp_path, monkeypatch
):
    """Each page gets DIFFERENT text keyed on `page.number`, so a reversed
    order, a dropped first or last page, or a shifted index each produce a
    different, catchable wrong order -- unlike identical text on every
    page, which cannot tell any of those bugs apart from a correct one."""
    monkeypatch.setattr(conversion, "get_settings", lambda: _ocr_settings())
    monkeypatch.setattr(
        conversion,
        "_ocr_page",
        lambda page, **kwargs: f"Article {page.number + 1}: texte de la page.",
    )

    result = conversion.convert_file(_scanned_pdf_pages(tmp_path / "scan3.pdf", 3))

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.page_count == 3
    positions = [result.markdown.find(f"Article {n}") for n in (1, 2, 3)]
    assert all(position != -1 for position in positions), positions
    assert positions[0] < positions[1] < positions[2], positions


def test_ocr_page_cap_exceeded_reports_skipped_with_the_limit_named(
    tmp_path, monkeypatch
):
    """F-16's cost guard: a scan longer than `ocr_max_pages` is Skipped,
    naming both the file's page count and the configured limit, and OCR is
    never even attempted -- the expensive part is what the cap exists to
    avoid."""
    monkeypatch.setattr(conversion, "get_settings", lambda: _ocr_settings(ocr_max_pages=2))

    def must_not_run(*args, **kwargs):
        raise AssertionError("_ocr_page must not run over the page cap")

    monkeypatch.setattr(conversion, "_ocr_page", must_not_run)

    result = conversion.convert_file(_scanned_pdf_pages(tmp_path / "scan3.pdf", 3))

    assert result.outcome is ConversionOutcome.SKIPPED
    assert result.outcome is not ConversionOutcome.FAILED
    assert "3" in result.reason
    assert "2" in result.reason


def test_ocr_finding_nothing_reports_skipped_with_the_ocr_specific_reason(
    tmp_path, monkeypatch
):
    """OCR being ATTEMPTED and finding nothing is a different fact from OCR
    never having run (the plain V1 wording), so it gets its own reason --
    asserted by presence of the new wording AND absence of the old one, so
    a mutation that reuses the V1 string cannot pass by accident."""
    monkeypatch.setattr(conversion, "get_settings", lambda: _ocr_settings())
    monkeypatch.setattr(conversion, "_ocr_page", lambda page, **kwargs: "")

    result = conversion.convert_file(_scanned_pdf(tmp_path / "scan.pdf"))

    assert result.outcome is ConversionOutcome.SKIPPED
    assert "OCR" in result.reason
    assert "no text layer" not in result.reason


def test_ocr_raising_reports_failed_and_never_stops_the_batch(
    tmp_path, monkeypatch, caplog
):
    """F-02 criterion 3, applied to the new failure mode: OCR crashing on
    one scanned file (a real risk -- it is third-party code fed rendered
    pixels) costs that one file, exactly like a converter crash already
    does, and never the rest of the sync."""
    monkeypatch.setattr(conversion, "get_settings", lambda: _ocr_settings())

    def explode(page, **kwargs):
        raise RuntimeError("tesseract blew up")

    monkeypatch.setattr(conversion, "_ocr_page", explode)

    scanned = _scanned_pdf(tmp_path / "scan.pdf")
    good = tmp_path / "fine.txt"
    good.write_text("Un texte tout a fait normal.", encoding="utf-8")

    with caplog.at_level(logging.ERROR, logger="conversion"):
        results = conversion.convert_files([scanned, good])

    by_name = {result.file_name: result for result in results}
    assert by_name["scan.pdf"].outcome is ConversionOutcome.FAILED
    assert "tesseract blew up" in by_name["scan.pdf"].reason
    assert by_name["fine.txt"].outcome is ConversionOutcome.CONVERTED


# --- OCR (F-16): real tessdata, skipped where it is absent (CI) --------
#
# This machine's tessdata_fast files are an external asset (DECISIONS.md),
# not committed to the repo, at a fixed local path. CI has none, so these
# skip there by design -- the plumbing above is what CI proves, and these
# are what a human ran locally as the actual proof OCR reads real pixels.

_TESSDATA_DIR = "C:/Users/lenovo/tessdata"
_HAS_TESSDATA = Path(_TESSDATA_DIR).is_dir()
_NO_TESSDATA_REASON = (
    "real tessdata_fast language files not found at C:/Users/lenovo/tessdata "
    "on this machine; they are an external asset, not part of the repo "
    "(see README), so CI and any other machine skip these by design"
)


def _real_ocr_settings(languages):
    return Settings(_env_file=None, ocr_tessdata_dir=_TESSDATA_DIR, ocr_languages=languages)


def _rasterized_text_pdf(path, texts, *, fontsize=24, dpi=200):
    """A PDF that genuinely has no text layer: each page is a bitmap of
    rendered text, built by writing real text onto a throwaway page,
    rendering it to a pixmap, and inserting that pixmap as a picture into
    a fresh page -- the exact recipe used to prove the OCR call works
    before this feature was implemented."""
    doc = pymupdf.open()
    for text in texts:
        source = pymupdf.open()
        source_page = source.new_page()
        source_page.insert_text((30, 60), text, fontsize=fontsize)
        pixmap = source_page.get_pixmap(dpi=dpi)
        source.close()
        page = doc.new_page()
        page.insert_image(page.rect, pixmap=pixmap)
    doc.save(path)
    doc.close()
    return path


_ARABIC_FONT_CSS = (
    "@font-face {font-family: arf; src: url(arial.ttf);} "
    "* {font-family: arf; direction: rtl; font-size: 24px;}"
)


def _rasterized_arabic_pdf(path, text, *, dpi=200):
    """Same recipe as `_rasterized_text_pdf`, but the source text is drawn
    with `insert_htmlbox` and a system font that actually has Arabic
    glyphs (`arial.ttf` does not by default carry them via `insert_text`,
    which is Latin-only)."""
    styled = pymupdf.open()
    styled_page = styled.new_page()
    styled_page.insert_htmlbox(
        pymupdf.Rect(30, 30, 500, 150),
        text,
        css=_ARABIC_FONT_CSS,
        archive=pymupdf.Archive("C:/Windows/Fonts"),
    )
    pixmap = styled_page.get_pixmap(dpi=dpi)
    styled.close()
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, pixmap=pixmap)
    doc.save(path)
    doc.close()
    return path


@pytest.mark.skipif(not _HAS_TESSDATA, reason=_NO_TESSDATA_REASON)
def test_real_ocr_finds_french_text_in_a_rasterized_scan(tmp_path, monkeypatch):
    settings = _real_ocr_settings("fra")
    monkeypatch.setattr(conversion, "get_settings", lambda: settings)
    path = _rasterized_text_pdf(
        tmp_path / "scan.pdf",
        ["Article 12: Le salarie a droit a un conge annuel paye."],
    )

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert "Article 12" in result.markdown
    assert result.page_count == 1


@pytest.mark.skipif(not _HAS_TESSDATA, reason=_NO_TESSDATA_REASON)
def test_real_ocr_maps_pages_in_order_not_off_by_one(tmp_path, monkeypatch):
    settings = _real_ocr_settings("fra")
    monkeypatch.setattr(conversion, "get_settings", lambda: settings)
    texts = [
        "Article 1: premier texte du document.",
        "Article 2: deuxieme texte totalement different.",
        "Article 3: troisieme et dernier texte.",
    ]
    path = _rasterized_text_pdf(tmp_path / "scan3.pdf", texts)

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert result.page_count == 3
    positions = [result.markdown.find(f"Article {n}") for n in (1, 2, 3)]
    assert all(position != -1 for position in positions), positions
    assert positions[0] < positions[1] < positions[2], positions


@pytest.mark.skipif(not _HAS_TESSDATA, reason=_NO_TESSDATA_REASON)
def test_real_ocr_finds_arabic_words(tmp_path, monkeypatch):
    """tessdata_fast's Arabic model is not perfect (recorded honestly in
    DECISIONS.md): this asserts the words it actually returned when this
    test was written, not a hoped-for exact transcription."""
    settings = _real_ocr_settings("ara")
    monkeypatch.setattr(conversion, "get_settings", lambda: settings)
    path = _rasterized_arabic_pdf(
        tmp_path / "scan_ar.pdf",
        "\u0627\u0644\u0645\u0627\u062f\u0629 12: \u0644\u0644\u0623\u062c\u064a\u0631 "
        "\u0627\u0644\u062d\u0642 \u0641\u064a \u0625\u062c\u0627\u0632\u0629 "
        "\u0633\u0646\u0648\u064a\u0629 \u0645\u062f\u0641\u0648\u0639\u0629 "
        "\u0627\u0644\u0623\u062c\u0631.",
    )

    result = conversion.convert_file(path)

    assert result.outcome is ConversionOutcome.CONVERTED
    assert "\u0627\u0644\u0645\u0627\u062f\u0629" in result.markdown  # "Article/clause"
    assert "\u0627\u0644\u0623\u062c\u0631" in result.markdown  # "the pay"


# --- batch behaviour: PRD F-02 criterion 3 ----------------------------


def test_one_bad_file_never_stops_the_others(tmp_path):
    """The exit gate's headline promise, exercised on a mixed corpus with
    the failure deliberately placed FIRST, so a batch that aborts on it
    loses everything after it.

    Every outcome is asserted, not just the count of survivors: asserting
    only "4 converted" would pass if the corrupted file had quietly
    converted into garbage too."""
    corrupt = tmp_path / "1-broken.pdf"
    corrupt.write_bytes(b"%PDF-1.7 not really\x00\xff")
    scan = _scanned_pdf(tmp_path / "2-scan.pdf")
    good_pdf = _pdf_with_headings(tmp_path / "3-good.pdf")
    good_docx = _docx(tmp_path / "4-note.docx")
    good_md = tmp_path / "5-readme.md"
    good_md.write_text("# Titre\n\nCorps.\n", encoding="utf-8")
    good_txt = tmp_path / "6-notes.txt"
    good_txt.write_text("Du texte.\n", encoding="utf-8")

    results = conversion.convert_files(
        [corrupt, scan, good_pdf, good_docx, good_md, good_txt]
    )

    assert [r.outcome for r in results] == [
        ConversionOutcome.FAILED,
        ConversionOutcome.SKIPPED,
        ConversionOutcome.CONVERTED,
        ConversionOutcome.CONVERTED,
        ConversionOutcome.CONVERTED,
        ConversionOutcome.CONVERTED,
    ]
    assert [r.file_name for r in results] == [
        "1-broken.pdf",
        "2-scan.pdf",
        "3-good.pdf",
        "4-note.docx",
        "5-readme.md",
        "6-notes.txt",
    ]
    assert all(r.markdown for r in results if r.converted)


def test_results_never_carry_a_contradictory_pair(tmp_path):
    """The invariant behind the dataclass docstring: markdown is set on
    CONVERTED only, reason on the other two only. Checked across the whole
    corpus in one place so no single rung can drift out of it."""
    corrupt = tmp_path / "broken.pdf"
    corrupt.write_bytes(b"%PDF-1.7 not really\x00\xff")
    paths = [
        corrupt,
        _scanned_pdf(tmp_path / "scan.pdf"),
        _pdf_with_headings(tmp_path / "good.pdf"),
        _docx(tmp_path / "note.docx"),
    ]

    for result in conversion.convert_files(paths):
        if result.converted:
            assert result.markdown and result.reason is None
        else:
            assert result.reason and result.markdown is None


def test_every_supported_extension_has_a_rung():
    """The ladder and the V1 set are one decision in two files. Adding
    "pptx" to `supported_document_extensions` without writing its
    converter would otherwise ship a silent Skipped on every deck the
    scan hands over.

    Asserted as equality, both directions: a missing rung fails, and so
    does a rung for a type the scan never yields."""
    assert set(conversion._CONVERTERS) == set(get_settings().supported_document_extensions)


@pytest.mark.parametrize(
    "name", ["doc.PDF", "note.DocX", "deck.PPTX", "readme.MD", "notes.TXT"]
)
def test_extension_matching_is_case_insensitive(tmp_path, name):
    """Windows and macOS hand over `REPORT.PDF` routinely. Dispatch shares
    `change_detection.file_type` precisely so a file the scan accepted can
    never be a file the ladder refuses."""
    path = tmp_path / name
    if path.suffix.lower() == ".pdf":
        _pdf_with_headings(path)
    elif path.suffix.lower() == ".docx":
        _docx(path)
    elif path.suffix.lower() == ".pptx":
        _pptx(path, ["Some text."])
    else:
        path.write_text("# Titre\n\nCorps.\n", encoding="utf-8")

    assert conversion.convert_file(path).outcome is ConversionOutcome.CONVERTED
