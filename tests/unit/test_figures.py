"""Figures: which pages are analysed, what is kept, where it is stored.

The fixture is a generated maintenance manual with a known answer key:
a vector diagram (boxes, lines, a circle, labels ON TOP of the shapes) on
page 1, a photo on page 2, a logo repeated in the header of every page,
and a framed warning paragraph that must never come out as an image.
"""

from __future__ import annotations

import pymupdf
import pytest

import figures
import sync
from chunking import Parent
from config import get_settings
from ui.conversation import _cards_for
from vector_store import SearchHit

_LABELS = ["Réservoir", "Pompe P-201", "Échangeur", "Compresseur"]
_PARA = (
    "La pompe P-201 alimente le circuit de refroidissement du compresseur. "
    "Avant toute intervention, l'opérateur ferme la vanne V-12."
)


def _settings(tmp_path, **overrides):
    return get_settings().model_copy(
        update={
            "figures_enabled": True,
            "figure_store_path": str(tmp_path / "figures"),
            "figure_explanations": "off",
            **overrides,
        }
    )


@pytest.fixture
def enabled(tmp_path, monkeypatch):
    settings = _settings(tmp_path)
    monkeypatch.setattr(figures, "get_settings", lambda: settings)
    return settings


def _manual(path, *, diagram=True, photo=True, text_only_rules=False, watermark=False):
    logo = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 30), False)
    logo.set_rect(logo.irect, (200, 30, 30))
    picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 400, 260), False)
    for x in range(400):
        for y in range(0, 260, 4):
            picture.set_pixel(x, y, (x % 255, (y * 3) % 255, 120))
    doc = pymupdf.open()
    for number in range(3):
        page = doc.new_page(width=595, height=842)
        page.insert_image(pymupdf.Rect(40, 20, 100, 50), pixmap=logo)
        if watermark:
            page.insert_image(pymupdf.Rect(70, 100, 525, 745), pixmap=picture, overlay=False)
        page.insert_text((60, 90), f"{number + 1}. Circuit de refroidissement", fontsize=16)
        page.insert_textbox(pymupdf.Rect(60, 110, 535, 200), _PARA, fontsize=11)
        if text_only_rules:
            for y in range(220, 700, 20):
                page.draw_line((60, y), (535, y))
        if number == 0 and diagram:
            shape = page.new_shape()
            boxes = [pymupdf.Rect(80 + i * 115, 250, 170 + i * 115, 310) for i in range(4)]
            for i, box in enumerate(boxes):
                shape.draw_rect(box)
                shape.finish(color=(0, 0, 0), fill=(0.85, 0.9, 1))
                if i < 3:
                    shape.draw_line((box.x1, 280), (box.x1 + 25, 280))
                    shape.finish(color=(0, 0, 0), width=2)
            shape.draw_circle((300, 380), 30)
            shape.finish(color=(0, 0, 0), fill=(1, 0.9, 0.8))
            shape.commit()
            for box, label in zip(boxes, _LABELS, strict=True):
                page.insert_textbox(box + (5, 20, -5, 0), label, fontsize=9, align=1)
            page.insert_text((278, 384), "V-12", fontsize=10)
            page.insert_text((60, 440), "Figure 1 : Schéma du circuit de refroidissement")
            frame = pymupdf.Rect(60, 470, 535, 560)
            page.draw_rect(frame, color=(0, 0, 0))
            page.insert_textbox(
                frame + (8, 8, -8, -8), "Attention : ne jamais ouvrir la vanne V-12.", fontsize=11
            )
        if number == 1 and photo:
            page.insert_image(pymupdf.Rect(97, 220, 497, 480), pixmap=picture)
            page.insert_text((60, 500), "Figure 2 : Photo de la pompe P-201 installée")
    doc.save(path)
    return path


# --- the cheap page gate --------------------------------------------------


def test_only_pages_with_a_real_figure_are_sent_to_the_layout_model(tmp_path, enabled):
    pages, repeated = figures._pdf_pages_with_visuals(_manual(tmp_path / "m.pdf"))

    assert pages == [1, 2]
    assert len(repeated) == 1, "the header logo is on every page, so it is decoration"


def test_a_logo_on_every_page_does_not_make_a_page_worth_analysing(tmp_path, enabled):
    pages, _ = figures._pdf_pages_with_visuals(
        _manual(tmp_path / "m.pdf", diagram=False, photo=False)
    )

    assert pages == []


def test_a_large_image_behind_every_page_is_a_watermark_not_a_figure(tmp_path, enabled):
    """The Moroccan Labour Code case: one background image over most of
    every page. Big enough to pass the size rule, so only the repeat rule
    keeps the whole law away from the layout model."""
    pages, repeated = figures._pdf_pages_with_visuals(
        _manual(tmp_path / "m.pdf", diagram=False, photo=False, watermark=True)
    )

    assert pages == []
    assert len(repeated) == 2, "the logo and the watermark"


def test_ruled_lines_under_text_are_not_a_diagram(tmp_path, enabled):
    """Underlines and table rules: horizontal lines only, no shapes."""
    pages, _ = figures._pdf_pages_with_visuals(
        _manual(tmp_path / "m.pdf", diagram=False, photo=False, text_only_rules=True)
    )

    assert pages == []


def test_a_single_frame_around_text_is_not_graphics(tmp_path, enabled):
    pdf = pymupdf.open(_manual(tmp_path / "m.pdf"))
    page = pdf[0]

    frame = pymupdf.Rect(60, 470, 535, 560)
    diagram = pymupdf.Rect(78, 248, 517, 412)
    assert not figures._has_graphics(page, frame, frozenset())
    assert figures._has_graphics(page, diagram, frozenset())


def test_figures_are_off_when_disabled(tmp_path, monkeypatch):
    settings = _settings(tmp_path, figures_enabled=False)
    monkeypatch.setattr(figures, "get_settings", lambda: settings)

    assert figures.extract_figures(_manual(tmp_path / "m.pdf")) == []


def test_a_broken_file_gives_no_figures_and_no_error(tmp_path, enabled):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.7 not really")

    assert figures.extract_figures(broken) == []


# --- the real layout model, once ------------------------------------------


def test_the_real_extractor_keeps_whole_figures_and_nothing_else(tmp_path, enabled):
    """The answer key, run through Docling for real (about a minute).

    Checks every accuracy promise at once: the diagram is ONE figure whose
    crop holds all its labels, the photo is kept, the logo and the framed
    paragraph are not figures, and each figure knows its caption, section,
    page and the text around it."""
    path = _manual(tmp_path / "manual.pdf")

    found = figures.extract_figures(path)

    assert [item.figure.page for item in found] == [1, 2]
    diagram, photo = (item.figure for item in found)
    assert diagram.caption.startswith("Figure 1")
    assert photo.caption.startswith("Figure 2")
    assert diagram.heading.startswith("1.")
    assert "P-201" in diagram.context_before
    # The crop must hold the whole diagram: every label's words inside it.
    pdf = pymupdf.open(path)
    words = {w[4]: pymupdf.Rect(w[:4]) for w in pdf[0].get_text("words")}
    ratio = 150 / 72
    box = pymupdf.Rect(0, 0, diagram.width_px / ratio, diagram.height_px / ratio)
    assert box.width >= (words["Compresseur"].x1 - words["Réservoir"].x0) - 2
    assert box.height >= (words["V-12"].y1 - words["Pompe"].y0) - 2


# --- storage --------------------------------------------------------------


def _extracted(index=0, source="m.pdf"):
    figure = figures.Figure(
        id=figures.figure_id(source, index, "ab" * 32),
        source_file=source,
        index=index,
        page=1,
        caption="Figure 1 : Schéma",
        heading="1. Circuit",
        context_before="avant",
        context_after="après",
        image_sha256="ab" * 32,
        width_px=200,
        height_px=100,
        explanation="Un schéma.",
    )
    return figures.ExtractedFigure(figure=figure, png=b"\x89PNG fake")


def test_saved_figures_load_back_and_are_deleted_with_their_file(tmp_path, enabled):
    item = _extracted()
    other = _extracted(source="other.pdf")
    figures.save_figures(workspace_id="ws1", figures=[item, other])

    assert figures.load_figure(workspace_id="ws1", figure_id=item.figure.id) == item.figure
    assert figures.figure_png_path(workspace_id="ws1", figure_id=item.figure.id).is_file()

    assert figures.delete_figures(workspace_id="ws1", source_file="m.pdf") == 1
    assert figures.load_figure(workspace_id="ws1", figure_id=item.figure.id) is None
    assert figures.figure_png_path(workspace_id="ws1", figure_id=item.figure.id) is None
    assert figures.load_figure(workspace_id="ws1", figure_id=other.figure.id) is not None


def test_workspace_delete_removes_every_figure(tmp_path, enabled):
    figures.save_figures(workspace_id="ws1", figures=[_extracted(0), _extracted(1)])

    assert figures.delete_workspace_figures(workspace_id="ws1") == 2
    assert figures.figures_of(workspace_id="ws1", source_file="m.pdf") == []


@pytest.mark.parametrize("bad", ["../x", "a/b", "", "..", "x.png"])
def test_an_unsafe_id_never_reaches_the_filesystem(tmp_path, enabled, bad):
    with pytest.raises(figures.UnsafeFigureIdError):
        figures.figure_png_path(workspace_id="ws1", figure_id=bad)


# --- attaching a figure to its text section -------------------------------


def test_a_figure_card_hangs_off_the_section_that_holds_its_caption():
    parents = [
        Parent(id="p1", text="Introduction générale du manuel.", source_file="m.pdf"),
        Parent(
            id="p2",
            text="Le circuit.\nFigure 1 : Schéma du circuit de refroidissement\nSuite.",
            source_file="m.pdf",
            section_label="1. Circuit",
        ),
    ]
    item = _extracted()
    item = figures.ExtractedFigure(
        figure=figures.Figure(
            **{**item.figure.__dict__, "caption": "Figure 1 : Schéma du circuit de refroidissement"}
        ),
        png=item.png,
    )

    (card,) = sync._figure_cards([item], parents)

    assert card.parent_id == "p2"
    assert card.section_label == "1. Circuit"
    assert card.figure_id == item.figure.id
    assert "Description générée" in card.text


def test_a_figure_with_no_place_in_the_text_goes_to_the_first_section():
    parents = [Parent(id="p1", text="Rien à voir.", source_file="m.pdf")]

    (card,) = sync._figure_cards([_extracted()], parents)

    assert card.parent_id == "p1"


# --- source cards ---------------------------------------------------------


def test_a_cited_figure_appears_on_its_source_card(tmp_path, enabled, monkeypatch):
    from agent.state import Source

    item = _extracted()
    figures.save_figures(workspace_id="ws1", figures=[item])
    section = "Le circuit de refroidissement. Détails."
    hits = [
        SearchHit("p2", "m.pdf", "1. Circuit", "Le circuit de refroidissement.", 0.9),
        SearchHit("p2", "m.pdf", "1. Circuit", item.figure.card_text(), 0.8, item.figure.id),
    ]

    (card,) = _cards_for([Source("m.pdf", "1. Circuit")], hits, {"p2": section}, "ws1")

    assert [ref.figure_id for ref in card.figures] == [item.figure.id]
    assert card.figures[0].explanation == "Un schéma."
    assert card.passages[0].highlighted, "the text hit still highlights the section"


def test_a_figure_whose_record_is_gone_is_left_off_the_card(tmp_path, enabled):
    from agent.state import Source

    hits = [SearchHit("p2", "m.pdf", None, "card", 0.8, "0" * 32)]

    (card,) = _cards_for([Source("m.pdf", None)], hits, {"p2": "texte"}, "ws1")

    assert card.figures == ()


# --- the description step, with a fake vision model ------------------------


def test_the_description_gets_the_image_and_the_context_and_is_trimmed(monkeypatch):
    from agent import vision

    seen = {}

    def fake_ask(model, system, user, png):
        seen.update(system=system, user=user, png=png)
        return "  Un schéma du circuit.  "

    monkeypatch.setattr(vision, "_vision_model", lambda: object())
    monkeypatch.setattr(vision, "_ask", fake_ask)

    text = vision.describe_figure(
        b"PNG", workspace="Maintenance", document="m.pdf", figure=_extracted().figure
    )

    assert text == "Un schéma du circuit."
    assert seen["png"] == b"PNG"
    assert "Maintenance" in seen["user"] and "Figure 1 : Schéma" in seen["user"]
    assert "avant" in seen["user"] and "après" in seen["user"]
    assert "Never invent" in seen["system"]


def test_a_failed_description_leaves_the_figure_without_one(monkeypatch):
    from agent import vision

    def broken(*_args):
        raise TimeoutError("provider timed out")

    monkeypatch.setattr(vision, "_vision_model", lambda: object())
    monkeypatch.setattr(vision, "_ask", broken)

    figure = _extracted().figure
    assert vision.describe_figure(b"PNG", workspace="", document="m.pdf", figure=figure) == ""


def test_no_description_when_descriptions_are_off(monkeypatch):
    from agent import vision

    settings = get_settings().model_copy(update={"figure_explanations": "off"})
    monkeypatch.setattr(vision, "get_settings", lambda: settings)

    assert vision._vision_model() is None


def test_strict_local_mode_without_a_vision_model_sends_nothing(monkeypatch):
    from agent import vision

    settings = get_settings().model_copy(
        update={
            "figure_explanations": "model",
            "model_mode": "strict_local",
            "vision_model_local": "",
        }
    )
    monkeypatch.setattr(vision, "get_settings", lambda: settings)

    assert vision._vision_model() is None


# --- PowerPoint and Word, through the real layout model --------------------


def _png(width, height):
    picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    for x in range(width):
        for y in range(0, height, 4):
            picture.set_pixel(x, y, (x % 255, (y * 3) % 255, 120))
    return picture.tobytes("png")


def test_a_slide_picture_is_kept_with_its_slide_caption_and_title(tmp_path, enabled):
    """A logo on every slide is dropped; the one real picture keeps its
    slide number, its "Figure 1" line as caption, and the slide title as
    heading -- and borrows no text from the next slide."""
    import io

    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for number in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Maintenance {number + 1}"
        slide.shapes.add_picture(
            io.BytesIO(_png(60, 30)), Inches(0.2), Inches(0.1), Inches(0.6), Inches(0.3)
        )
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(1))
        body.text_frame.text = "La pompe P-201 alimente le circuit de refroidissement."
        if number == 1:
            slide.shapes.add_picture(
                io.BytesIO(_png(400, 260)), Inches(2), Inches(2.5), Inches(5), Inches(3.2)
            )
            caption = slide.shapes.add_textbox(Inches(2), Inches(5.8), Inches(6), Inches(0.5))
            caption.text_frame.text = "Figure 1 : Vue de la pompe P-201"
    path = tmp_path / "deck.pptx"
    deck.save(path)

    (item,) = figures.extract_figures(path)

    assert item.figure.page == 2
    assert item.figure.caption == "Figure 1 : Vue de la pompe P-201"
    assert item.figure.heading == "Maintenance 2"
    assert "P-201" in item.figure.context_before
    assert "Maintenance 3" not in item.figure.context_after


def test_a_word_picture_keeps_its_caption_and_the_paragraphs_around_it(tmp_path, enabled):
    import io

    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading("1. Circuit de refroidissement", 1)
    document.add_paragraph("La pompe P-201 alimente le circuit.")
    document.add_picture(io.BytesIO(_png(400, 260)), width=Inches(4))
    document.add_paragraph("Figure 1 : Vue de la pompe P-201")
    document.add_paragraph("Après la figure, on ferme la vanne V-12.")
    path = tmp_path / "doc.docx"
    document.save(path)

    (item,) = figures.extract_figures(path)

    assert item.figure.page is None
    assert item.figure.caption == "Figure 1 : Vue de la pompe P-201"
    assert item.figure.heading == "1. Circuit de refroidissement"
    assert item.figure.context_before == "La pompe P-201 alimente le circuit."
    assert item.figure.context_after == "Après la figure, on ferme la vanne V-12."
